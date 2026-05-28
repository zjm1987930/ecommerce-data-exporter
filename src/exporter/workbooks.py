from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border


DELAYED_TAB_ORDER = [
    "延迟发货",
    "缺货",
    "虚假发货/虚假轨迹",
]

DELAYED_TAB_HEADERS = {
    "延迟发货": [
        "订单编号",
        "商品编号",
        "补偿金额",
        "补偿时间",
        "承诺发货时间",
        "发货时间",
        "成团时间",
        "揽收时间",
        "补偿原因",
        "商品名称",
        "商品类型",
    ],
    "缺货": [
        "订单编号",
        "商品编号",
        "商品名称",
        "成团时间",
        "承诺发货时间",
        "报名价",
        "补偿时间",
        "补偿金额",
    ],
    "虚假发货/虚假轨迹": [
        "订单编号",
        "商品ID",
        "商品名称",
        "发货时间",
        "违规类型",
        "实收金额",
        "商品报名价之和",
        "补偿金额",
        "补偿时间",
        "快递单号",
    ],
}


def _detect_tab_name(path: Path) -> str | None:
    name = path.name
    if "延迟发货" in name:
        return "延迟发货"
    if "缺货" in name:
        return "缺货"
    if "虚假发货" in name or "虚假轨迹" in name:
        return "虚假发货/虚假轨迹"
    return None


def _safe_sheet_name(name: str) -> str:
    return re.sub(r"[:\\/?*\[\]]", "_", name)[:31]


def merge_delayed_shipping(raw_dir: Path, output_path: Path) -> Path:
    files_by_tab: dict[str, Path] = {}
    for path in raw_dir.glob("*.xlsx"):
        tab = _detect_tab_name(path)
        if tab:
            files_by_tab[tab] = path
    return merge_delayed_shipping_files(files_by_tab, output_path)


def merge_delayed_shipping_files(files_by_tab: dict[str, Path], output_path: Path) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for tab in DELAYED_TAB_ORDER:
        ws_out = workbook.create_sheet(_safe_sheet_name(tab))
        source = files_by_tab.get(tab)
        if not source:
            ws_out.append(DELAYED_TAB_HEADERS[tab])
            continue
        wb_in = load_workbook(source, data_only=False)
        ws_in = wb_in[wb_in.sheetnames[0]]
        for row in ws_in.iter_rows(values_only=True):
            ws_out.append(list(row))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


NEGATIVE_EXPERIENCE_HEADERS = [
    "订单编号",
    "商品信息",
    "问题类型",
    "补偿金额（元）",
    "扣款进度",
    "创建时间",
    "处理状态",
    "操作",
]

NEGATIVE_EXPERIENCE_SHEET_ORDER = [
    "拼多多金喆",
    "拼多多逸观",
    "拼多多大雁",
    "拼多多大雁母婴玩具",
    "拼多多大雁母婴",
    "拼多多大雁书籍",
    "拼多多大雁童书",
    "拼多多逸观母婴玩具",
    "拼多多逸观书籍",
]


@dataclass(frozen=True)
class NegativeExperienceRecord:
    order_no: str
    product_name: str
    product_id: str
    problem_type: str
    compensation_amount: str
    deduction_progress: str
    created_at: str
    status: str
    operation: str
    suggestion: str = ""


def write_negative_experience_workbook(
    records_by_sheet: dict[str, list[NegativeExperienceRecord]],
    output_path: Path,
) -> Path:
    if output_path.exists():
        workbook = load_workbook(output_path)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)

    ordered_names = list(NEGATIVE_EXPERIENCE_SHEET_ORDER)
    for name in records_by_sheet:
        if name not in ordered_names:
            ordered_names.append(name)

    for sheet_name in ordered_names:
        safe_name = _safe_sheet_name(sheet_name)
        if safe_name in workbook.sheetnames:
            if sheet_name not in records_by_sheet:
                continue
            index = workbook.sheetnames.index(safe_name)
            del workbook[safe_name]
            ws = workbook.create_sheet(safe_name, index)
        else:
            ws = workbook.create_sheet(safe_name)
        _write_negative_experience_sheet(ws, records_by_sheet.get(sheet_name, []))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _write_negative_experience_sheet(ws, records: list[NegativeExperienceRecord]) -> None:
    ws.append(NEGATIVE_EXPERIENCE_HEADERS)
    widths = {
        "A": 24,
        "B": 42,
        "C": 24,
        "D": 16,
        "E": 16,
        "F": 22,
        "G": 24,
        "H": 14,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    header_fill = PatternFill("solid", fgColor="F5F7FA")
    thin = Side(style="thin", color="E5E7EB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 24

    row = 2
    for record in records:
        ws.cell(row=row, column=1, value=record.order_no)
        ws.cell(row=row, column=3, value=record.problem_type)
        ws.cell(row=row, column=4, value=record.compensation_amount)
        ws.cell(row=row, column=5, value=record.deduction_progress)
        ws.cell(row=row, column=6, value=record.created_at)
        ws.cell(row=row, column=7, value=record.status)
        ws.cell(row=row, column=8, value=record.operation)

        ws.cell(row=row + 1, column=2, value=record.product_name)
        ws.cell(row=row + 2, column=2, value=record.product_id)
        if record.suggestion:
            ws.cell(row=row + 1, column=3, value=record.suggestion)

        for column in [1, 4, 5, 6, 7, 8]:
            ws.merge_cells(start_row=row, start_column=column, end_row=row + 2, end_column=column)
        if not record.suggestion:
            ws.merge_cells(start_row=row, start_column=3, end_row=row + 2, end_column=3)

        for row_index in range(row, row + 3):
            ws.row_dimensions[row_index].height = 34 if row_index != row + 1 else 72
            for column in range(1, 9):
                cell = ws.cell(row=row_index, column=column)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
        row += 3
